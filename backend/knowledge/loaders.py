"""
文档加载器 —— 直接使用解析库 + LangChain Document 包装。
===========================================================================
LangChain v1.x 中 langchain-community 不再维护，官方推荐直接使用
解析库并将结果包装为 langchain_core.documents.Document。

每种文件类型用独立解析库提取文本，统一返回 list[Document]：
    - page_content: 文本内容
    - metadata: { page_number, source }

支持的文件类型：
    PDF:  PyMuPDF (fitz)，按页提取，保留 page_number
    DOCX: python-docx，按段落提取，page_number=-1
    TXT:  直接读取，自动处理编码，page_number=-1

仍兼容 LangChain Document 流转：loaders → chunker → vector_store
===========================================================================
"""
import logging
from pathlib import Path

from langchain_core.documents import Document

logger = logging.getLogger(__name__)


def load_document(file_path: Path, file_type: str) -> list[Document]:
    """
    根据文件类型加载文档，返回 LangChain Document 列表。

    Args:
        file_path:  文件路径
        file_type:  文件类型（pdf/docx/txt）

    Returns:
        LangChain Document 列表，含 page_content 和 metadata
    """
    file_type = file_type.lower().lstrip(".")

    if file_type == "pdf":
        return _load_pdf(file_path)
    elif file_type == "docx":
        return _load_docx(file_path)
    elif file_type == "txt":
        return _load_txt(file_path)
    else:
        raise ValueError(f"暂不支持的文件类型: {file_type}")


def _load_pdf(file_path: Path) -> list[Document]:
    """
    使用 PyMuPDF (fitz) 按页提取 PDF 文本。

    每页一个 Document，metadata 含 page_number。
    """
    import fitz  # PyMuPDF

    docs = []
    doc = fitz.open(file_path)
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text().strip()
        if text:
            docs.append(Document(
                page_content=text,
                metadata={
                    "page_number": page_num + 1,
                    "source": file_path.name,
                },
            ))
    doc.close()

    logger.info("PDF 加载: %s, %d 页有文本", file_path.name, len(docs))
    return docs


def _load_docx(file_path: Path) -> list[Document]:
    """
    使用 python-docx 提取段落文本。

    DOCX 不分页，page_number=-1。
    """
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    full_text = "\n".join(
        para.text.strip() for para in doc.paragraphs if para.text.strip()
    )

    if not full_text:
        return []

    docs = [Document(
        page_content=full_text,
        metadata={
            "page_number": -1,
            "source": file_path.name,
        },
    )]

    logger.info("DOCX 加载: %s", file_path.name)
    return docs


def _load_txt(file_path: Path) -> list[Document]:
    """
    直接读取 TXT 文件，自动处理编码。

    TXT 不分页，page_number=-1。
    """
    text = None
    for encoding in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            text = file_path.read_text(encoding=encoding).strip()
            if text:
                break
        except (UnicodeDecodeError, Exception):
            continue

    if not text:
        logger.warning("TXT 加载失败，无法识别编码: %s", file_path.name)
        return []

    docs = [Document(
        page_content=text,
        metadata={
            "page_number": -1,
            "source": file_path.name,
        },
    )]

    logger.info("TXT 加载: %s", file_path.name)
    return docs
